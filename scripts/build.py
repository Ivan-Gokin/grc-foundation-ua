#!/usr/bin/env python3
"""Build the site into _site/ for GitHub Pages. CI runs it on every push to main and publishes the result.

What it adds to the files kept in the repository (nothing generated is stored in git):
  data/provisions.csv, data/mappings.csv   from the workbooks in mappings/*.xlsx
  data/eli.jsonld                          the register and chains as ELI metadata (JSON-LD):
                                           ELI ontology 1.5 and, for bills and EU proposals, ELI-DL
  data/grc-foundation-ua.xlsx              one readable workbook for people: register, chains and
                                           mappings with names and texts joined in

ELI notes: every source is an eli:LegalResource (bills and EU proposals are eli:Work). Official ELIs
(EU acts) are absolute; the project's ELI-pattern identifiers are written relative to the file
("../eli/ua/..."), so they resolve against the published site. Links that start from a draft are
written as eli-dl:foresees_change_of, except eli-dl:adopts. Titles: eli:title on a LegalExpression in
the official language; this project's translations go to rdfs:label only.

Usage:  python3 scripts/build.py            # then: python3 -m http.server 8000 -d _site
"""
import csv
import datetime as dt
import json
import shutil
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

import mapdata

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
OUT = ROOT / "_site"
COPY = ["index.html", "404.html", "assets", "content", "data"]

ELI_PROPS = {"based_on", "amends", "repeals", "transposes", "applies", "ensures_implementation_of",
             "commences", "cites", "related_to"}
DL_PROPS = {"adopts"}
IN_FORCE = {"in_force": "eli:InForce-inForce", "repealed": "eli:InForce-notInForce",
            "not_yet_in_force": "eli:InForce-notInForce"}
UKR = "http://publications.europa.eu/resource/authority/country/UKR"
LANG = {"uk": "http://publications.europa.eu/resource/authority/language/UKR",
        "en": "http://publications.europa.eu/resource/authority/language/ENG"}
APPROX_DATE_TYPES = {"nd_tzi", "dstu", "vst", "nato_policy"}  # adoption date known only to the year

CONTEXT = {
    "eli": "http://data.europa.eu/eli/ontology#",
    "eli-dl": "http://data.europa.eu/eli/eli-draft-legislation-ontology#",
    "rdfs": "http://www.w3.org/2000/01/rdf-schema#",
    "skos": "http://www.w3.org/2004/02/skos/core#",
    "dct": "http://purl.org/dc/terms/",
    "owl": "http://www.w3.org/2002/07/owl#",
    "xsd": "http://www.w3.org/2001/XMLSchema#",
    "date_document": {"@id": "eli:date_document", "@type": "xsd:date"},
    **{p: {"@id": f"eli:{p}", "@type": "@id"} for p in sorted(ELI_PROPS | {"in_force", "jurisdiction", "type_document",
                                                                          "passed_by", "language", "is_realized_by"})},
    "adopts": {"@id": "eli-dl:adopts", "@type": "@id"},
    "foresees_change_of": {"@id": "eli-dl:foresees_change_of", "@type": "@id"},
    "id_local": "eli:id_local",
    "number": "eli:number",
    "title": "eli:title",
    "title_short": "eli:title_short",
    "label": "rdfs:label",
    "seeAlso": {"@id": "rdfs:seeAlso", "@type": "@id"},
    "prefLabel": "skos:prefLabel",
}


def iri(eli):
    return eli if eli.startswith("http") else "../" + eli


def lang_values(row, base):
    return [{"@value": row[f"{base}_{l}"], "@language": l} for l in ("uk", "en") if row.get(f"{base}_{l}")]


def eli_document(meta, sources, types, relations):
    """The register and its chains as ELI metadata (JSON-LD), see docs/BLUEPRINT.md 3.1 and 3.4."""
    by_id = {s["id"]: s for s in sources}
    links = {}
    for r in relations:
        src = by_id[r["source_id"]]
        prop = r["relation"]
        if src["status"] == "draft" and prop not in DL_PROPS:
            prop = "foresees_change_of"
        links.setdefault(r["source_id"], {}).setdefault(prop, []).append(iri(by_id[r["related_id"]]["eli"]))

    graph = [{
        "@id": "",
        "@type": "dct:Dataset",
        "dct:title": [{"@value": "Реєстр нормативних джерел GRC-фундаменту", "@language": "uk"},
                      {"@value": "GRC foundation register of legal sources", "@language": "en"}],
        "dct:conformsTo": [{"@id": "http://data.europa.eu/eli/ontology"}, {"@id": "http://data.europa.eu/eli/eli-draft-legislation-ontology"}],
        "owl:versionInfo": meta.get("version", ""),
        "dct:modified": {"@value": meta.get("date", ""), "@type": "xsd:date"},
        "rdfs:comment": [{"@value": "Ідентифікатори, що не починаються з http://data.europa.eu/eli/, - неофіційні: вони побудовані за шаблоном ELI цим проєктом, бо офіційного ELI в Україні поки немає.", "@language": "uk"},
                         {"@value": "Identifiers that do not start with http://data.europa.eu/eli/ are unofficial: this project builds them on the ELI URI template because Ukraine has no official ELI yet.", "@language": "en"}],
    }]
    for t in sorted(types):
        graph.append({"@id": f"#type-{t}", "@type": "skos:Concept", "prefLabel": lang_values(types[t], "label")})
        if types[t].get("issuer_uk"):
            graph.append({"@id": f"#issuer-{t}", "@type": "dct:Agent", "label": lang_values(types[t], "issuer")})

    for s in sorted(sources, key=lambda x: x["eli"]):
        draft = s["status"] == "draft"
        node = {
            "@id": iri(s["eli"]),
            "@type": "eli:Work" if draft else "eli:LegalResource",
            "id_local": s["id"],
            "type_document": f"#type-{s['type']}",
        }
        if s["number"]:
            node["number"] = s["number"]
        if s["adopted"] and not (s["type"] in APPROX_DATE_TYPES and s["adopted"].endswith("-01-01")):
            node["date_document"] = s["adopted"]
        if s["status"] in IN_FORCE:
            node["in_force"] = IN_FORCE[s["status"]]
        if s["id"].startswith("ua-"):
            node["jurisdiction"] = UKR
        if types.get(s["type"], {}).get("issuer_uk"):
            node["passed_by"] = f"#issuer-{s['type']}"
        official = "en" if s["scope"] == "reference" or s["type"] == "nato_policy" else "uk"
        if not s["eli"].startswith("http"):  # EUR-Lex publishes the official expressions of EU acts itself
            title = s[f"title_{official}"] or s[f"short_ref_{official}"]
            expr = {"@id": f"{iri(s['eli'])}/{'ukr' if official == 'uk' else 'eng'}", "@type": "eli:LegalExpression",
                    "language": LANG[official], "title": {"@value": title, "@language": official}}
            if s[f"short_ref_{official}"] and s[f"short_ref_{official}"] != title:
                expr["title_short"] = {"@value": s[f"short_ref_{official}"], "@language": official}
            node["is_realized_by"] = [expr]
        node["label"] = lang_values(s, "title") or lang_values(s, "short_ref")
        if s["url"]:
            node["seeAlso"] = s["url"]
        for prop, targets in sorted(links.get(s["id"], {}).items()):
            node[prop] = sorted(targets)
        graph.append(node)
    return {"@context": CONTEXT, "@graph": graph}


# ---------------------------------------------------------------- readable workbook
REL_UK = {"equal": "= збігається із заходом", "subset_of": "⊂ вужче за захід", "superset_of": "⊃ ширше за захід",
          "intersects_with": "∩ перетинається із заходом", "no_relationship": "∅ немає відповідника"}
LINK_UK = {"based_on": "видано на підставі", "amends": "змінює", "repeals": "скасовує", "transposes": "транспонує",
           "applies": "узгоджено з", "ensures_implementation_of": "забезпечує виконання", "commences": "вводить у дію",
           "cites": "посилається на", "related_to": "пов'язаний з", "adopts": "ухвалено за проєктом"}
STATUS_UK = {"in_force": "чинний", "repealed": "втратив чинність", "not_yet_in_force": "ще не чинний", "draft": "проєкт",
             "unknown": "невідомо"}
REVIEW_UK = {"ai_draft": "ШІ, не вичитано", "draft": "чернетка", "imported": "імпорт", "reviewed": "перевірено"}
COV_UK = {"full": "повне", "partial": "часткове", "none": "немає"}
HEAD = PatternFill("solid", fgColor="E8E9F6")


def _sheet(wb, title, header, rows, widths):
    ws = wb.create_sheet(title)
    ws.append(header)
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    for c in ws[1]:
        c.font, c.fill = Font(bold=True), HEAD
    for row in ws.iter_rows(min_row=1):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions


def readable_workbook(path, meta, sources, types, sectors, appl, relations, provisions, mappings, frameworks, controls):
    src = {s["id"]: s for s in sources}
    ref = lambda s: s["short_ref_uk"] or s["title_uk"] or s["id"]
    site = meta.get("site_url", "")
    eli = lambda s: s["eli"] if s["eli"].startswith("http") else site + s["eli"]  # full IRI on the published site
    tags = {}
    for a in appl:
        label = sectors[a["sector_id"]]["label_uk"]
        reqs = [sectors[x]["label_uk"] for x in (a.get("requires") or "").split(";") if x in sectors]
        tags.setdefault(a["source_id"], []).append(" + ".join([label] + reqs))
    wb = openpyxl.Workbook()
    about = wb.active
    about.title = "Про файл"
    ai = sum(m["review_status"] == "ai_draft" for m in mappings)
    lines = [
        ["GRC-фундамент для України: реєстр, ланцюги та зіставлення / GRC foundation for Ukraine"],
        [f"Версія даних {meta.get('version', '')} від {meta.get('date', '')}. Згенеровано {dt.date.today().isoformat()}. Сайт: {meta.get('site_url', '')}"],
        [],
        [f"Увага: {ai} зіставлень згенеровано ШІ і ще не вичитано людиною (статус «ШІ, не вичитано»). Звіряйте з текстом акта."],
        ["Це довідковий матеріал, а не юридична консультація. / A reference, not legal advice."],
        [],
        ["Аркуші:"],
        ["Реєстр - нормативні джерела з описом, статусом, типами суб'єктів і ідентифікатором ELI."],
        ["Ланцюги - зв'язки між актами (властивості онтології ELI) з правовою підставою."],
    ] + [[f"{fw['name']} {fw['version']} - заходи стандарту і положення актів, що їх вимагають (NIST IR 8477); «—» - вимога без відповідника."] for fw in frameworks] + [
        [],
        ["ELI, що не починаються з http://data.europa.eu/eli/, неофіційні: їх побудовано за шаблоном ELI в цьому проєкті, бо офіційного ELI в Україні поки немає. Вони відкривають акт у реєстрі на сайті."],
        ["Ліцензія даних: CC BY 4.0. Тексти ISO/IEC 27001 не відтворено: лише ідентифікатори та власні короткі назви."],
        ["Машинні дані: CSV (UTF-8) і eli.jsonld у тій самій теці data/, опис стовпців у datapackage.json."],
    ]
    for l in lines:
        about.append(l)
    about.column_dimensions["A"].width = 120
    about["A1"].font = Font(bold=True, size=13)

    _sheet(wb, "Реєстр",
           ["№", "Перелік", "Вид", "Акт", "Назва", "Дата ухвалення", "Статус", "Обов'язковість", "Опис", "Кого стосується",
            "Типи суб'єктів", "ELI", "Офіційний текст", "Перевірено", "id"],
           [[int(s["register_no"]), "Законодавство України" if s["scope"] == "core" else "Довідково: право ЄС", types[s["type"]]["label_uk"],
             ref(s), s["title_uk"], s["adopted"], STATUS_UK.get(s["status"], s["status"]),
             "обов'язковий" if s["binding"] == "mandatory" else "рекомендаційний", s["description_uk"], s["concerns_uk"],
             "; ".join(tags.get(s["id"], [])), eli(s), s["url"], s["last_verified"], s["id"]] for s in sorted(sources, key=lambda s: int(s["register_no"]))],
           [6, 14, 18, 28, 50, 12, 12, 14, 60, 40, 30, 34, 30, 11, 22])
    _sheet(wb, "Ланцюги",
           ["Акт", "Зв'язок", "Властивість ELI", "Пов'язаний акт", "Правова підстава", "Де знайдено", "Достовірність", "ELI акта", "ELI пов'язаного"],
           [[ref(src[r["source_id"]]) + (" (проєкт)" if src[r["source_id"]]["status"] == "draft" else ""),
             LINK_UK.get(r["relation"], r["relation"]) + (" (за проєктом)" if src[r["source_id"]]["status"] == "draft" and r["relation"] != "adopts" else ""),
             ("eli-dl:" if r["relation"] == "adopts" else "eli:") + r["relation"], ref(src[r["related_id"]]), r["basis_uk"], r["source_url"],
             {"high": "висока", "medium": "середня"}.get(r["confidence"], r["confidence"]), eli(src[r["source_id"]]), eli(src[r["related_id"]])]
            for r in relations],
           [32, 22, 26, 32, 60, 30, 12, 34, 34])
    prov = {p["id"]: p for p in provisions}
    for fw in frameworks:
        ctl = controls[fw["id"]]
        name = lambda c: (c.get("text_uk") or c.get("name_uk") or c.get("name_en", "")) if fw["id"] == "csf-2.0" else (c.get("name_uk") or c.get("name_en", ""))
        rows = []
        order = {c["id"]: i for i, c in enumerate(sorted(ctl.values(), key=lambda c: float(c["sort"] or 0)))}
        for m in sorted([m for m in mappings if m["framework_id"] == fw["id"]], key=lambda m: (order.get(m["control_id"], 10**6), m["provision_id"])):
            p = prov[m["provision_id"]]
            s = src[p["source_id"]]
            c = ctl.get(m["control_id"], {})
            rows.append([m["control_id"] or "—", name(c) if c else "немає відповідника", ref(s), p["ref_uk"], p["summary_uk"],
                         REL_UK.get(m["relationship"], m["relationship"]), COV_UK.get(m["coverage"], m["coverage"]), m["parameter_uk"],
                         m["policy_uk"], m["notes_uk"], p["notes"],
                         "; ".join(sectors[x]["label_uk"] for x in p["applies_to"].split(";") if x in sectors),
                         REVIEW_UK.get(m["review_status"], m["review_status"]), m["mapped_by"], m["reviewed_by"], m["date"], eli(s)])
        _sheet(wb, {"iso27001-2022": "ISO 27001", "csf-2.0": "NIST CSF 2.0"}.get(fw["id"], fw["id"])[:31],
               ["Захід", "Назва заходу", "Акт", "Пункт", "Вимога", "Відношення", "Покриття", "Параметр за законом", "Куди внести",
                "Примітка", "Застереження", "Стосується лише", "Статус", "Хто зіставив", "Хто перевірив", "Дата", "ELI акта"],
               rows, [9, 40, 28, 24, 60, 22, 10, 24, 34, 50, 34, 24, 15, 22, 16, 11, 34])
    wb.save(path)


# ---------------------------------------------------------------- build
def read_csv(name):
    with open(DATA / name, encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


def write_csv(path, fields, rows):
    with open(path, "w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=fields, lineterminator="\n", extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


def bust_cache(page):
    """Add ?v=<content hash> to the page's own scripts and stylesheets, so a browser that cached the previous
    version (GitHub Pages lets browsers keep files for 10 minutes) fetches the new ones as soon as it gets the page."""
    import hashlib
    import re
    html = page.read_text(encoding="utf-8")

    def stamp(m):
        f = OUT / m.group(2)
        v = hashlib.sha256(f.read_bytes()).hexdigest()[:10] if f.exists() else ""
        return f'{m.group(1)}="{m.group(2)}?v={v}"' if v else m.group(0)
    page.write_text(re.sub(r'(src|href)="(assets/[^"?]+\.(?:js|css))"', stamp, html), encoding="utf-8")


def main():
    if OUT.exists():
        shutil.rmtree(OUT)
    OUT.mkdir()
    for name in COPY:
        src = ROOT / name
        (shutil.copytree if src.is_dir() else shutil.copy2)(src, OUT / name)

    bust_cache(OUT / "index.html")
    meta = json.loads((DATA / "meta.json").read_text(encoding="utf-8"))
    sources, relations = read_csv("sources.csv"), read_csv("source_relations.csv")
    types = {r["id"]: r for r in read_csv("source_types.csv")}
    sectors = {r["id"]: r for r in read_csv("sectors.csv")}
    frameworks = [f for f in read_csv("frameworks.csv") if f["tier"] == "organisation"]
    controls = {f["id"]: {c["id"]: c for c in read_csv(f["controls_file"])} for f in frameworks}

    provisions, mappings, errors = mapdata.read_all({f["id"] for f in frameworks})
    if errors:
        raise SystemExit("mappings/*.xlsx have errors, run scripts/validate.py:\n" + "\n".join(errors[:20]))
    write_csv(OUT / "data" / "provisions.csv", mapdata.PROVISION_FIELDS, provisions)
    write_csv(OUT / "data" / "mappings.csv", mapdata.MAPPING_FIELDS, mappings)

    doc = eli_document(meta, sources, types, relations)
    (OUT / "data" / "eli.jsonld").write_text(json.dumps(doc, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")

    readable_workbook(OUT / "data" / "grc-foundation-ua.xlsx", meta, sources, types, sectors, read_csv("applicability.csv"),
                      relations, provisions, mappings, frameworks, controls)
    print(f"built {OUT.relative_to(ROOT)}/: {len(provisions)} provisions, {len(mappings)} mappings, {len(sources)} sources")


if __name__ == "__main__":
    main()
