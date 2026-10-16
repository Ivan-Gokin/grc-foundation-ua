#!/usr/bin/env python3
"""One-off import (2026-10-16) of two AI-generated STRM workbooks (NIST IR 8477) into the data model.

  ISO27001_2022_x_NBU_STRM_NIST_IR_8477.xlsx   ISO/IEC 27001:2022 <-> 8 NBU resolutions
  STRM_NIST_CSF_2.0_vs_ZU_80-94.xlsx           NIST CSF 2.0 <-> Law of Ukraine No. 80/94-VR

Both workbooks were generated with AI and have NOT been proofread by a person. Every imported
mapping therefore gets review_status=ai_draft; it becomes `reviewed` only after a person has checked
it against the text of the act (see docs/BLUEPRINT.md, 4.4).

What the script does:
  - one provision per (act, element) from the STRM sheets, plus the "unmapped requirements" sheets
    (requirements with no counterpart in the framework, mapped with relationship=no_relationship
    and an empty control_id);
  - converts the relationship direction: the NBU workbook reads "ISO element [rel] act", the data
    model reads "provision [rel] control", so "superset of" becomes subset_of;
  - adds the policy_uk/policy_en columns to mappings.csv ("Цільова політика" / "Куди включити");
  - marks the earlier 22 draft mappings written by Claude as ai_draft as well;
  - tags Law 80/94 as applying to every organisation (universal norms) and to CII owners, as the
    workbook's applicability sheet says.

Kept for provenance; data/*.csv are the source of truth afterwards. Do not re-run it on edited data:
it replaces every row it created. Run it on the pre-import data only.

Usage: python3 scripts/import_strm_2026_10.py <nbu.xlsx> <zu.xlsx>
"""
import csv
import re
import sys
from pathlib import Path

import openpyxl

DATA = Path(__file__).resolve().parent.parent / "data"
DATE = "2026-10-16"

NBU_ACTS = {
    "П-95": ("ua-nbu-95-2017", "Текст: первісна редакція 2017 р. (дзеркало iplex.com.ua)."),
    "П-178": ("ua-nbu-178-2022", "Текст: редакція зі змінами № 24/2025, 69/2025, 66/2026."),
    "П-43": ("ua-nbu-43-2021", "Текст: редакція від 01.08.2022; зміни Постанови № 66/2026 (п. 9-10 розділу II) не звірено."),
    "П-4": ("ua-nbu-4-2021", "Зміни Постанови № 24/2025 - за переказом, нумерація додатків може відрізнятися."),
    "П-64": ("ua-nbu-64-2018", "Нумерацію розділу VI (операційний ризик, TPSP) звірити з чинною редакцією від 01.07.2026."),
    "П-88": ("ua-nbu-88-2019", "Нумерація за первісною редакцією 2019 р.; звірити з чинною (від 26.06.2026)."),
    "П-99": ("ua-nbu-99-2025", "Посилання на рівні розділів; текст за проєктом НБУ від 20.12.2024, звірити з офіційним."),
    "П-67": ("ua-nbu-67-2024", ""),
}
ZU_ID = "ua-law-80-94-vr"
ZU_NOTE = "Нумерацію частин ст. 8-10 звірити з чинним текстом (ред. Закону № 4336-IX)."
AI_NOTE = "Виділено ШІ, не вичитано."
MAPPED_BY_NBU = "AI (STRM workbook ISO27001_2022_x_NBU)"
MAPPED_BY_ZU = "AI (STRM workbook NIST_CSF_2.0_vs_ZU_80-94)"
MAPPED_BY_CLAUDE = "AI (Claude)"

REL_NBU = {"superset of": "subset_of", "intersects with": "intersects_with", "equal": "equal", "subset of": "superset_of"}
REL_ZU = {"subset of": "subset_of", "intersects with": "intersects_with", "equal": "equal", "superset of": "superset_of"}
COVERAGE = {"subset_of": "full", "equal": "full", "intersects_with": "partial", "superset_of": "partial"}

# Law 80/94 addressee codes -> applies_to. Д = systems with state information resources or
# restricted information (shown for state-ir and the public sector); Ф = financial sector.
ZU_CII_PARTS = {"ст. 8 ч. 1", "ст. 8 ч. 2", "ст. 8 ч. 3", "ст. 8 ч. 7", "ст. 8 ч. 11", "ст. 8 ч. 12"}


def addressee(code, ref):
    code = code.strip()
    if code.startswith("У"):
        return ""
    if code.startswith("Ф"):
        return "finance"
    if "держтаємн" in code:
        return "state-secrets"
    if code.startswith("Держоргани"):
        return "public-sector"
    out = ["state-ir", "public-sector"]
    if ref in ZU_CII_PARTS:
        out.append("kii")
    return ";".join(out)


ROMAN = {"I": 1, "II": 2, "III": 3, "IV": 4, "V": 5, "VI": 6, "VII": 7, "VIII": 8, "IX": 9, "X": 10,
         "XI": 11, "XII": 12, "XIII": 13}


def ref_en(uk):
    """Best-effort English rendering of a Ukrainian paragraph reference; '' if anything is left over."""
    s = uk
    for a, b in [
        (r"\(у ред\. Постанови № (\d+) від ([\d.]+)\)", r"(as amended by Resolution No. \1 of \2)"),
        (r"\(у ред\. № (\d+)/(\d+)\)", r"(as amended by No. \1/\2)"),
        (r"\(не звірено\)", "(not verified)"),
        (r"\(ланцюгові хмарні послуги\)", "(chained cloud services)"),
        (r"\(вимоги до [Дд]оговору\)", "(contract requirements)"),
        (r"\(інформування НБУ\)", "(notifying the NBU)"),
        (r"\(управління ризиками під час використання хмарних послуг\)", "(risk management when using cloud services)"),
        (r"\(набрання чинності, перехідні положення\)", "(entry into force, transitional provisions)"),
        (r"\(графік\)", "(schedule)"),
        (r"\(абз\. (\d+)\)", r"(subpara. \1)"),
        (r"абз\. (\d+)", r"subpara. \1"),
        (r"розділів ([IVX]+)[–-]([IVX]+)", r"of Sections \1–\2"),
        (r"розділів ([IVX]+), ([IVX]+)[–-]([IVX]+)", r"of Sections \1, \2–\3"),
        (r"розділів ([IVX]+), ([IVX]+)", r"of Sections \1, \2"),
        (r"розділу ([IVX]+)[–-]([IVX]+)", r"of Sections \1–\2"),
        (r"розділу ([IVX]+), ([IVX]+)", r"of Sections \1, \2"),
        (r"розділів ([IVX]+)", r"of Section \1"),
        (r"розділу ([IVX]+)", r"of Section \1"),
        (r"розділи ([IVX]+)[–-]([IVX]+)", r"Sections \1–\2"),
        (r"розділ ([IVX]+)", r"Section \1"),
        (r"глави ([\d-]+)[–-]([\d-]+)", r"Chapters \1–\2"),
        (r"глав ([\d]+)[–-]([\d]+)", r"of Chapters \1–\2"),
        (r"глави ([\d-]+)", r"of Chapter \1"),
        (r"глава (\d+)", r"Chapter \1"),
        (r"додатки (\d+)[–-](\d+)", r"Annexes \1–\2"),
        (r"додаток \(п\. ", "Annex (para. "),
        (r"додаток", "Annex"),
        (r"постанова", "Resolution"),
        (r"ст\. ", "Art. "),
        (r" ч\. ", " part "),
        (r"п\. ", "para. "),
    ]:
        s = re.sub(a, b, s)
    s = s.replace("Chapters 44–47 of Section VI", "Chapters 44–47 of Section VI")
    if re.search(r"[А-Яа-яІіЇїЄєҐґ]", s):
        return ""
    return s


def zu_slug(ref):
    s = ref.replace("ст. ", "a").replace(" ч. ", "-p").replace(" п. ", "-i").replace("; ", ".")
    s = s.replace("–", "-").replace(" ", "")
    return s


def sheet_rows(wb, title, skip):
    ws = wb[title]
    rows = []
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        if i < skip:
            continue
        r = ["" if c is None else str(c).strip() for c in r]
        if any(r):
            rows.append(r)
    return rows


def read_csv(name):
    with open(DATA / name, encoding="utf-8", newline="") as fh:
        rd = csv.DictReader(fh)
        return rd.fieldnames, list(rd)


def write_csv(name, fields, rows):
    with open(DATA / name, "w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=fields, lineterminator="\n")
        w.writeheader()
        for r in rows:
            w.writerow({k: r.get(k, "") for k in fields})


def clean_policy(s):
    s = (s or "").strip()
    s = re.sub(r"^Аркуш «Реквізити»;?\s*", "", s)
    return "" if s in ("—", "-") else s


def is_definitional(ref, text):
    return ref.startswith("постанова") or re.match(r"^(Сфера|Поширюється|Положення не встановлює|Чинність|Набрання чинності)", text) is not None


def main():
    nbu_path, zu_path = sys.argv[1], sys.argv[2]
    p_fields, provisions = read_csv("provisions.csv")
    m_fields, mappings = read_csv("mappings.csv")
    a_fields, appl = read_csv("applicability.csv")

    ours = {MAPPED_BY_NBU, MAPPED_BY_ZU}
    new_sources = {v[0] for v in NBU_ACTS.values()} | {ZU_ID}
    provisions = [p for p in provisions if p["source_id"] not in new_sources]
    mappings = [m for m in mappings if m["mapped_by"] not in ours]

    # earlier Claude drafts are AI as well
    for m in mappings:
        if m["mapped_by"] in ("draft (Claude)", MAPPED_BY_CLAUDE):
            m["mapped_by"] = MAPPED_BY_CLAUDE
            m["review_status"] = "ai_draft"
    for p in provisions:
        if p["source_id"] == "ua-cmu-373-2006" and p["notes"].startswith("paragraph numbers"):
            p["notes"] = ("Номери пунктів - за редакцією до постанови КМУ № 1531/2025 (нова редакція Мінімальних вимог); "
                          "звірити з чинним текстом. / Paragraph numbers are from the edition before CMU Resolution 1531/2025; "
                          "verify against the current text.")

    if "policy_uk" not in m_fields:
        i = m_fields.index("notes_en") + 1
        m_fields[i:i] = ["policy_uk", "policy_en"]

    new_p, new_m = [], []

    # ---------------------------------------------------------------- NBU x ISO 27001
    wb = openpyxl.load_workbook(nbu_path, data_only=True)
    prov_key, counters = {}, {}

    def nbu_provision(code, ref, text, gap=False):
        src, caveat = NBU_ACTS[code]
        key = (src, ref, "gap" if gap else "")
        if key in prov_key:
            return prov_key[key], False
        counters[src] = counters.get(src, 0) + 1
        pid = f"{src}:el-{counters[src]:03d}"
        prov_key[key] = pid
        new_p.append({"id": pid, "source_id": src, "ref_uk": ref, "ref_en": ref_en(ref), "summary_uk": text, "summary_en": "",
                      "provision_type": "definitional" if gap and is_definitional(ref, text) else "substantive",
                      "applies_to": "", "notes": " ".join(x for x in (AI_NOTE, caveat) if x)})
        return pid, True

    summary_of = {}
    for r in sheet_rows(wb, "STRM", 1):
        ctl, _, rationale, rel, _, code, ref, text, strength, notes, policy = r[:11]
        if code not in NBU_ACTS:  # "Усі 8 актів": not-related rows for A.7.6, A.7.7, A.8.6 have no provision
            continue
        pid, created = nbu_provision(code, ref, text)
        if created:
            summary_of[pid] = text
        relationship = REL_NBU[rel]
        note = notes if text == summary_of[pid] else f"{text} {notes}".strip()
        new_m.append({"provision_id": pid, "framework_id": "iso27001-2022", "control_id": ctl, "relationship": relationship,
                      "rationale": rationale.lower(), "strength": strength if strength not in ("", "0") else "",
                      "coverage": COVERAGE[relationship], "notes_uk": note, "policy_uk": clean_policy(policy),
                      "mapped_by": MAPPED_BY_NBU, "review_status": "ai_draft", "date": DATE})
    for r in sheet_rows(wb, "Незамаплені вимоги акта", 2):
        code, ref, text, _, policy = r[:5]
        pid, _ = nbu_provision(code, ref, text, gap=True)
        new_m.append({"provision_id": pid, "framework_id": "iso27001-2022", "control_id": "", "relationship": "no_relationship",
                      "rationale": "semantic", "coverage": "none", "policy_uk": clean_policy(policy),
                      "mapped_by": MAPPED_BY_NBU, "review_status": "ai_draft", "date": DATE})

    # ---------------------------------------------------------------- Law 80/94 x CSF 2.0
    wb = openpyxl.load_workbook(zu_path, data_only=True)
    zu_ids = {}

    def zu_provision(ref, text, code, gap=False):
        key = (ref, gap)
        if key in zu_ids:
            return zu_ids[key]
        pid = f"{ZU_ID}:{zu_slug(ref)}" + ("-gap" if gap and (ref, False) in zu_ids else "")
        zu_ids[key] = pid
        new_p.append({"id": pid, "source_id": ZU_ID, "ref_uk": ref, "ref_en": ref_en(ref), "summary_uk": text, "summary_en": "",
                      "provision_type": "substantive", "applies_to": addressee(code, ref), "notes": f"{AI_NOTE} {ZU_NOTE}"})
        return pid

    for r in sheet_rows(wb, "STRM", 1):
        _, ctl, _, rel, rationale, indirect, ref, text, explain, code, policy = r[:11]
        pid = zu_provision(ref, text, code)
        note = ("Опосередкований зв'язок: прямої норми немає, результат випливає з обов'язку Закону. " if indirect == "так" else "") + explain
        relationship = REL_ZU[rel]
        new_m.append({"provision_id": pid, "framework_id": "csf-2.0", "control_id": ctl, "relationship": relationship,
                      "rationale": rationale.lower(), "coverage": COVERAGE[relationship], "notes_uk": note.strip(),
                      "policy_uk": clean_policy(policy), "mapped_by": MAPPED_BY_ZU, "review_status": "ai_draft", "date": DATE})
    for r in sheet_rows(wb, "Незамаплені вимоги", 1):
        ref, text, code, policy = r[:4]
        pid = zu_provision(ref, text, code, gap=True)
        new_m.append({"provision_id": pid, "framework_id": "csf-2.0", "control_id": "", "relationship": "no_relationship",
                      "rationale": "semantic", "coverage": "none", "policy_uk": clean_policy(policy),
                      "mapped_by": MAPPED_BY_ZU, "review_status": "ai_draft", "date": DATE})

    # ---------------------------------------------------------------- applicability of Law 80/94
    have = {(a["source_id"], a["sector_id"]) for a in appl}
    for sector, basis in [("general", "ст. 4-7, ст. 8 ч. 1, ст. 9 ч. 1: базові норми для будь-якого власника системи і володільця інформації (робоча книга STRM, аркуш «Застосовність», ШІ)"),
                          ("kii", "ст. 8 ч. 1-3, 7, 11-12: власники і розпорядники ОКІІ будь-якої форми власності (робоча книга STRM, аркуш «Застосовність», ШІ)")]:
        if (ZU_ID, sector) not in have:
            appl.append({"source_id": ZU_ID, "sector_id": sector, "basis": basis, "verified": "no"})

    write_csv("provisions.csv", p_fields, provisions + new_p)
    write_csv("mappings.csv", m_fields, mappings + new_m)
    write_csv("applicability.csv", a_fields, appl)
    print(f"provisions +{len(new_p)}, mappings +{len(new_m)}")


if __name__ == "__main__":
    main()
