"""Read and write the mapping workbooks in mappings/*.xlsx.

Each workbook holds a group of acts (for example all NBU resolutions):
  - sheet "Вимоги": one row per provision (requirement) of an act;
  - one sheet per framework, named by its id in data/frameworks.csv ("iso27001-2022", "csf-2.0"):
    one row per link between a provision and a control, NIST IR 8477 style.
A link points to its provision by act + paragraph ("Акт" + "Пункт"), so the requirement text lives
in one place only. Control "—" means: checked against the whole framework, no counterpart.

read_all() returns the provisions and mappings in the flat shape the site and the checks use
(the columns of the former provisions.csv and mappings.csv), plus a list of errors.
"""
import datetime as dt
import re
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
MAPPINGS = ROOT / "mappings"
PROV_SHEET = "Вимоги"
NONE_CONTROL = {"—", "–", "-", "немає", "none"}

TYPE = {"substantive": "обов'язок", "delegating": "делегує", "definitional": "визначення"}
COVERAGE = {"full": "повне", "partial": "часткове", "none": "немає"}
STATUS = {"ai_draft": "ШІ (не вичитано)", "draft": "чернетка", "imported": "імпорт", "reviewed": "перевірено"}
RELATIONSHIP = {"equal": "equal", "subset_of": "subset of", "superset_of": "superset of",
                "intersects_with": "intersects with", "no_relationship": "no relationship"}
RATIONALE = ["semantic", "syntactic", "functional"]
DEFAULT_COVERAGE = {"equal": "full", "subset_of": "full", "superset_of": "partial", "intersects_with": "partial",
                    "no_relationship": "none"}

# (field, header, width, help)
PROV_COLS = [
    ("source_id", "Акт", 20, "Ідентифікатор акта з data/sources.csv, наприклад ua-nbu-95-2017."),
    ("code", "Код", 14, "Короткий постійний код положення в межах акта (p40, a8-p2). Не змінюйте його після публікації. Якщо не заповнити, його буде утворено з пункту."),
    ("ref_uk", "Пункт", 32, "Посилання на пункт так, як його цитують: п. 40 розділу IV; ст. 8 ч. 2. Унікальне в межах акта."),
    ("summary_uk", "Вимога", 70, "Що вимагає положення, своїми словами."),
    ("provision_type", "Тип", 13, "обов'язок - зіставляється; делегує - лише доручає іншому акту; визначення - сфера чи терміни."),
    ("applies_to", "Стосується лише", 22, "Необов'язково: типи суб'єктів з data/sectors.csv через ';', якщо положення вужче за акт."),
    ("notes", "Застереження", 40, "Редакція, з якою працювали, що треба звірити."),
    ("ref_en", "Пункт (EN)", 24, "Необов'язково."),
    ("summary_en", "Вимога (EN)", 50, "Необов'язково."),
]
LINK_COLS = [
    ("source_id", "Акт", 20, "Ідентифікатор акта, як в аркуші Вимоги."),
    ("ref_uk", "Пункт", 32, "Пункт точно так, як в аркуші Вимоги (скопіюйте)."),
    ("control_id", "Захід", 11, "Ідентифікатор заходу стандарту (A.8.5, GV.OC-03). «—» - у стандарті немає відповідника."),
    ("relationship", "Відношення", 15, "NIST IR 8477, читається «положення [відношення] захід»: subset of - положення вужче за захід; superset of - ширше; intersects with - перетинаються; equal - збігаються."),
    ("rationale", "Обґрунтування", 13, "semantic - зміст; syntactic - формулювання; functional - результат."),
    ("strength", "Сила", 7, "Необов'язково, 1-10."),
    ("coverage", "Покриття", 11, "Необов'язково: якщо порожньо, subset of / equal = повне, інакше часткове."),
    ("parameter_uk", "Параметр за законом", 24, "Значення, яке закон встановлює для параметра заходу: «не рідше 1 разу на рік»."),
    ("policy_uk", "Куди внести", 34, "Внутрішній документ, куди варто внести вимогу."),
    ("notes_uk", "Примітка", 50, "Пояснення зв'язку."),
    ("review_status", "Статус", 17, "ШІ (не вичитано) / чернетка / імпорт / перевірено. «перевірено» ставить друга людина і вписує себе в «Хто перевірив»."),
    ("mapped_by", "Хто зіставив", 22, "Ім'я або «AI (...)» для згенерованих рядків."),
    ("reviewed_by", "Хто перевірив", 18, ""),
    ("date", "Дата", 11, "Дата останньої зміни рядка."),
    ("parameter_en", "Параметр (EN)", 20, "Необов'язково."),
    ("policy_en", "Куди внести (EN)", 20, "Необов'язково."),
    ("notes_en", "Примітка (EN)", 30, "Необов'язково."),
]
# column order of the generated data/provisions.csv and data/mappings.csv
PROVISION_FIELDS = ["id", "source_id", "ref_uk", "ref_en", "summary_uk", "summary_en", "provision_type", "applies_to", "notes"]
MAPPING_FIELDS = ["provision_id", "framework_id", "control_id", "relationship", "rationale", "strength", "coverage",
                  "parameter_uk", "parameter_en", "notes_uk", "notes_en", "policy_uk", "policy_en", "mapped_by",
                  "reviewed_by", "review_status", "date"]
LISTS = {"provision_type": TYPE, "coverage": COVERAGE, "review_status": STATUS, "relationship": RELATIONSHIP,
         "rationale": {r: r for r in RATIONALE}}

TRANSLIT = dict(zip("абвгґдеєжзиіїйклмнопрстуфхцчшщьюя", ["a", "b", "v", "h", "g", "d", "e", "ie", "zh", "z", "y", "i", "i", "i", "k", "l",
                    "m", "n", "o", "p", "r", "s", "t", "u", "f", "kh", "ts", "ch", "sh", "shch", "", "iu", "ia"]))


def code_from_ref(ref):
    s = "".join(TRANSLIT.get(c.lower(), c.lower()) for c in ref)
    s = s.replace("st.", "a").replace("ch.", "p").replace("p.", "p")
    return re.sub(r"[^a-z0-9.]+", "-", s).strip("-")[:40] or "x"


def _text(v):
    if v is None:
        return ""
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _from_label(field, value):
    if field not in LISTS or not value:
        return value, True
    v = value.strip()
    for code, label in LISTS[field].items():
        if v.lower() in (code, label.lower(), code.replace("_", " ")):
            return code, True
    if field == "relationship" and v.lower() == "not related to":
        return "no_relationship", True
    return v, False


def _rows(ws, cols, where, errors):
    header = [_text(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1))]
    by_header = {h: f for f, h, _, _ in cols}
    unknown = [h for h in header if h and h not in by_header]
    if unknown:
        errors.append(f"{where}: unknown columns {unknown}")
    idx = {by_header[h]: i for i, h in enumerate(header) if h in by_header}
    for missing in [h for f, h, _, _ in cols if f not in idx and f in ("source_id", "ref_uk")]:
        errors.append(f"{where}: column «{missing}» is missing")
    out = []
    for n, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        r = {f: _text(row[i]) if i < len(row) else "" for f, i in idx.items()}
        if not any(r.values()):
            continue
        for f in list(r):
            r[f], ok = _from_label(f, r[f])
            if not ok:
                errors.append(f"{where}, row {n}: «{r[f]}» is not a valid value for «{dict((c[0], c[1]) for c in cols)[f]}»")
        r["_where"] = f"{where}, row {n}"
        out.append(r)
    return out


def read_all(frameworks, directory=MAPPINGS):
    """frameworks: ids of the framework sheets. Returns (provisions, mappings, errors)."""
    provisions, mappings, errors = [], [], []
    seen_ref, seen_id = {}, {}
    for path in sorted(Path(directory).glob("*.xlsx")):
        if path.name.startswith("~$"):  # Excel lock file
            continue
        wb = openpyxl.load_workbook(path, data_only=True)
        name = path.name
        if PROV_SHEET not in wb.sheetnames:
            errors.append(f"{name}: sheet «{PROV_SHEET}» is missing")
            continue
        for extra in [s for s in wb.sheetnames if s != PROV_SHEET and s not in frameworks]:
            errors.append(f"{name}: sheet «{extra}» is neither «{PROV_SHEET}» nor a framework id {sorted(frameworks)}")
        local = {}
        for r in _rows(wb[PROV_SHEET], PROV_COLS, f"{name} / {PROV_SHEET}", errors):
            key = (r.get("source_id", ""), r.get("ref_uk", ""))
            if not all(key):
                errors.append(f"{r['_where']}: «Акт» and «Пункт» are required")
                continue
            if key in seen_ref:
                errors.append(f"{r['_where']}: {key[0]} «{key[1]}» is already in {seen_ref[key]}")
                continue
            code = r.get("code") or code_from_ref(key[1])
            pid = f"{key[0]}:{code}"
            if pid in seen_id:
                errors.append(f"{r['_where']}: code «{code}» of {key[0]} is already used in {seen_id[pid]}")
                continue
            seen_ref[key] = seen_id[pid] = r["_where"]
            p = {"id": pid, "source_id": key[0], "ref_uk": key[1], "ref_en": r.get("ref_en", ""),
                 "summary_uk": r.get("summary_uk", ""), "summary_en": r.get("summary_en", ""),
                 "provision_type": r.get("provision_type") or "substantive", "applies_to": r.get("applies_to", ""),
                 "notes": r.get("notes", ""), "_where": r["_where"]}
            provisions.append(p)
            local[key] = pid
        for fw in [s for s in wb.sheetnames if s in frameworks]:
            for r in _rows(wb[fw], LINK_COLS, f"{name} / {fw}", errors):
                key = (r.get("source_id", ""), r.get("ref_uk", ""))
                if key not in local:
                    errors.append(f"{r['_where']}: {key[0]} «{key[1]}» is not in the «{PROV_SHEET}» sheet of {name}")
                    continue
                control = r.get("control_id", "")
                if control in NONE_CONTROL:
                    control, r["relationship"] = "", "no_relationship"
                rel = r.get("relationship", "")
                m = {f: r.get(f, "") for f, _, _, _ in LINK_COLS if f not in ("source_id", "ref_uk")}
                m.update(provision_id=local[key], framework_id=fw, control_id=control,
                         coverage=r.get("coverage") or DEFAULT_COVERAGE.get(rel, ""),
                         review_status=r.get("review_status") or "draft", _where=r["_where"])
                mappings.append(m)
    return provisions, mappings, errors


# ---------------------------------------------------------------- writing (used to create workbooks)
HEAD_FILL = PatternFill("solid", fgColor="E8E9F6")


def _sheet(ws, cols, rows):
    ws.append([h for _, h, _, _ in cols])
    for i, (f, h, w, help_) in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font, cell.fill = Font(bold=True), HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if help_:
            cell.comment = Comment(help_, "grc-foundation-ua", width=300, height=110)
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    for r in rows:
        ws.append([LISTS.get(f, {}).get(r.get(f, ""), r.get(f, "")) if f in LISTS else r.get(f, "") for f, _, _, _ in cols])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    last = max(ws.max_row, 2) + 500
    for i, (f, _, _, _) in enumerate(cols, start=1):
        if f in LISTS:
            dv = DataValidation(type="list", formula1='"' + ",".join(LISTS[f].values()) + '"', allow_blank=True)
            letter = openpyxl.utils.get_column_letter(i)
            dv.add(f"{letter}2:{letter}{last}")
            ws.add_data_validation(dv)


def write_workbook(path, provisions, links_by_framework):
    """provisions: dicts with PROV_COLS fields; links_by_framework: {framework_id: [dicts with LINK_COLS fields]}."""
    wb = openpyxl.Workbook()
    _sheet(wb.active, PROV_COLS, provisions)
    wb.active.title = PROV_SHEET
    for fw, links in links_by_framework.items():
        _sheet(wb.create_sheet(fw), LINK_COLS, links)
    wb.save(path)
